#!/usr/bin/env python3
"""Ежедневный отчёт по метрикам сервера: VictoriaMetrics -> xlsx -> Telegram.

    report.py [--date YYYY-MM-DD] [--no-send] [--config /etc/srvmon/report.env] [--out DIR]

Без --date: сегодня, если рабочий день уже закончился (последняя точка прошла), иначе вчера
(так таймер с Persistent=true после простоя RPi выгружает пропущенный день).
"""
import argparse
import datetime as dt
import logging
import math
import os
import sys
import time
from collections import Counter, defaultdict
from zoneinfo import ZoneInfo

import requests
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

log = logging.getLogger("report")

STEP = 300  # 5 минут
RED = PatternFill("solid", fgColor="FFC7CE")
GREY = PatternFill("solid", fgColor="EDEDED")
HEAD = PatternFill("solid", fgColor="DDEBF7")

# Порядок интерфейсов: внешние, lo, docker0, bridge, veth, прочие
IFACE_ORDER = [("eth", 0), ("en", 0), ("lo", 1), ("docker", 2), ("br-", 3), ("veth", 4)]


def load_config(path):
    """KEY=VALUE файл; значения из окружения (systemd EnvironmentFile) имеют приоритет."""
    if path and os.path.exists(path):
        with open(path) as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#") and "=" in line:
                    k, v = line.split("=", 1)
                    v = v.strip()
                    if len(v) >= 2 and v[0] == v[-1] and v[0] in "'\"":
                        v = v[1:-1]
                    os.environ.setdefault(k.strip(), v)
    c = os.environ
    return {
        "vm_url": c.get("VM_URL", "http://127.0.0.1:8428"),
        "vm_auth": (c["VM_AUTH_USER"], c["VM_AUTH_PASSWORD"]) if c.get("VM_AUTH_USER") else None,
        "job": c.get("SERVER_JOB", "server-node"),
        "disk_re": c.get("DISK_REGEX", r"(sd[a-z]+|vd[a-z]+|nvme[0-9]+n[0-9]+|mmcblk[0-9]+)"),
        "tz": ZoneInfo(c.get("REPORT_TZ", "Asia/Tbilisi")),            # рабочий день и расписание
        "dtz": ZoneInfo(c.get("REPORT_DISPLAY_TZ", "Europe/Moscow")),  # время в Timestamp отчёта
        "start": c.get("WORK_START", "11:00"),
        "end": c.get("WORK_END", "20:00"),
        "wait_max": int(c.get("WAIT_MAX_SEC", "600")),
        "queue_url_re": c.get("VMAGENT_URL_RE", "1:.*"),
        "tg_token": c.get("TG_BOT_TOKEN", ""),
        "tg_chat": c.get("TG_CHAT_ID", ""),
        "hc_url": c.get("HC_DAILY_REPORT_URL", ""),
        "am_url": c.get("ALERTMANAGER_URL", "http://127.0.0.1:9093"),
        "out_dir": c.get("OUT_DIR", "/var/lib/srvmon/reports"),
    }


class VM:
    def __init__(self, cfg):
        self.url, self.auth = cfg["vm_url"].rstrip("/"), cfg["vm_auth"]

    def _get(self, path, **params):
        r = requests.get(self.url + path, params=params, auth=self.auth, timeout=60)
        r.raise_for_status()
        d = r.json()
        if d.get("status") != "success":
            raise RuntimeError(f"VM error: {d}")
        return d["data"]["result"]

    def range(self, q, start, end):
        return self._get("/api/v1/query_range", query=q, start=start, end=end, step=f"{STEP}s", nocache=1)

    def instant(self, q):
        return self._get("/api/v1/query", query=q, nocache=1)

    def push_metric(self, line):
        requests.post(self.url + "/api/v1/import/prometheus", data=line, auth=self.auth, timeout=30).raise_for_status()


def fval(s):
    try:
        return float(s)
    except (TypeError, ValueError):
        return float("nan")


def series_map(result, key=None):
    """{label_value or None: {ts: float}}"""
    out = {}
    for s in result:
        k = s["metric"].get(key) if key else None
        out[k] = {int(float(t)): fval(v) for t, v in s["values"]}
    return out


def iface_key(name):
    for prefix, prio in IFACE_ORDER:
        if name.startswith(prefix):
            return (prio, name)
    return (9, name)


class Column:
    def __init__(self, title, data, kind, resets=None, transient=False):
        self.title, self.data, self.kind = title, data, kind  # kind: net|cpu|ram|disk
        self.resets = resets or {}
        self.transient = transient  # может отсутствовать часть дня (veth, br-*)


def collect(vm, cfg, start, end):
    j, dre = cfg["job"], cfg["disk_re"]
    net = {}
    for direction, metric in (("in", "receive"), ("out", "transmit")):
        c = f'node_network_{metric}_bytes_total{{job="{j}"}}'
        net[direction] = (series_map(vm.range(f"rate({c}[5m]) * 8 / 1e6", start, end), "device"),
                          series_map(vm.range(f"resets({c}[5m])", start, end), "device"))
    cpu = series_map(vm.range(f'100 - avg(rate(node_cpu_seconds_total{{job="{j}",mode="idle"}}[5m])) * 100', start, end))
    cpu_rst = series_map(vm.range(f'sum(resets(node_cpu_seconds_total{{job="{j}",mode="idle"}}[5m]))', start, end))
    ram = series_map(vm.range(
        f'(1 - avg_over_time(node_memory_MemAvailable_bytes{{job="{j}"}}[5m])'
        f' / avg_over_time(node_memory_MemTotal_bytes{{job="{j}"}}[5m])) * 100', start, end))
    disk = {}
    for direction, metric in (("read", "reads"), ("write", "writes")):
        c = f'node_disk_{metric}_completed_total{{job="{j}",device=~"{dre}"}}'
        disk[direction] = (series_map(vm.range(f"rate({c}[5m])", start, end), "device"),
                           series_map(vm.range(f"resets({c}[5m])", start, end), "device"))

    cols = []
    ifaces = sorted(set(net["in"][0]) | set(net["out"][0]), key=iface_key)
    for i in ifaces:
        tr = i.startswith(("veth", "br-", "docker"))
        for d in ("in", "out"):
            cols.append(Column(f"Traffic {d} {i}", net[d][0].get(i, {}), "net", net[d][1].get(i, {}), tr))
    cols.append(Column("CPU load (%)", cpu.get(None, {}), "cpu", cpu_rst.get(None, {})))
    cols.append(Column("RAM used (%)", ram.get(None, {}), "ram"))
    for dev in sorted(set(disk["read"][0]) | set(disk["write"][0])):
        for d in ("read", "write"):
            cols.append(Column(f"Disk IOPS {d} {dev}", disk[d][0].get(dev, {}), "disk", disk[d][1].get(dev, {})))
    return cols


def validate(cols, expected, cutoff=None):
    """-> (issues[(ts, col, value, problem)], cell_status{(ts, title): (fill, text)}, complete_ts set)

    cutoff: точки позже этого момента ещё не наступили (окно усреднения неполное) — их значения
    отбрасываются, ячейки серые, в проблемы не попадают.
    """
    issues, cells = [], {}
    core = [c for c in cols if not c.transient]
    for c in cols:
        if cutoff is not None:
            c.data = {t: v for t, v in c.data.items() if t <= cutoff}
        present = sorted(t for t in c.data if t in expected)
        first, last = (present[0], present[-1]) if present else (None, None)
        for ts in expected:
            v = c.data.get(ts)
            prob = None
            if cutoff is not None and ts > cutoff:
                cells[(ts, c.title)] = (GREY, "ещё не наступило")
                continue
            if v is None:
                if c.transient and (first is None or ts < first or ts > last):
                    cells[(ts, c.title)] = (GREY, "интерфейса не было")
                    continue
                prob = "нет данных"
            elif math.isnan(v) or math.isinf(v):
                prob = "NaN/Inf"
            elif v < 0:
                prob = "отрицательное значение"
            elif c.kind in ("cpu", "ram") and v > 100:
                prob = "вне диапазона 0–100"
            elif c.kind == "ram" and v == 0:
                prob = "RAM = 0"
            if prob is None and c.resets.get(ts, 0) > 0:
                prob = "сброс счётчика"
            if prob:
                cells[(ts, c.title)] = (RED, prob)
                issues.append((ts, c.title, v, prob))
    complete = {ts for ts in expected if all(ts in c.data and not math.isnan(c.data[ts]) for c in core)}
    return issues, cells, complete


def buffer_state(vm, cfg):
    """(lag_sec, pending_bytes) — свежесть данных сервера и остаток буфера vmagent для этой RPi."""
    lag = vm.instant(f'time() - tlast_over_time(up{{job="{cfg["job"]}"}}[1h])')
    pend = vm.instant(f'max(vmagent_remotewrite_pending_data_bytes{{url=~"{cfg["queue_url_re"]}"}})')
    lag_s = fval(lag[0]["value"][1]) if lag else float("inf")
    pend_b = fval(pend[0]["value"][1]) if pend else float("nan")
    return lag_s, pend_b


def build_xlsx(path, day, cols, expected, cells, issues, complete, notes, tz, dtz):
    """tz — пояс рабочего дня, dtz — пояс, в котором выводятся Timestamp."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Данные"
    headers = ["Timestamp"] + [c.title for c in cols]
    ws.append(headers)
    for i, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True), HEAD
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for r, ts in enumerate(expected, 2):
        tcell = ws.cell(row=r, column=1, value=dt.datetime.fromtimestamp(ts, dtz).strftime("%Y-%m-%d %H:%M"))
        if ts not in complete:
            future = all(cells.get((ts, c.title), (None, ""))[1] == "ещё не наступило" for c in cols) if cols else False
            tcell.fill = GREY if future else RED
        for ci, c in enumerate(cols, 2):
            v = c.data.get(ts)
            cell = ws.cell(row=r, column=ci)
            if v is not None and not (math.isnan(v) or math.isinf(v)):
                cell.value = round(v, 2)
                cell.number_format = "0.00"
            elif v is not None:
                cell.value = str(v)
            st = cells.get((ts, c.title))
            if st:
                cell.fill = st[0]
                if st[0] is RED:
                    cell.comment = Comment(st[1], "report.py")
    ws.freeze_panes = "B2"
    ws.column_dimensions["A"].width = 17
    for i in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.row_dimensions[1].height = 45

    ch = wb.create_sheet("Проверка")
    future = {ts for ts in expected
              if cols and all(cells.get((ts, c.title), (None, ""))[1] == "ещё не наступило" for c in cols)}
    missing = [ts for ts in expected if ts not in complete and ts not in future]
    rows = [
        ("Дата", day.isoformat()),
        ("Интервал", f"{dt.datetime.fromtimestamp(expected[0], dtz):%H:%M} – {dt.datetime.fromtimestamp(expected[-1], dtz):%H:%M} {dtz.key} "
                     f"({dt.datetime.fromtimestamp(expected[0], tz):%H:%M} – {dt.datetime.fromtimestamp(expected[-1], tz):%H:%M} {tz.key}), "
                     "шаг 5 мин, среднее за предыдущие 5 мин"),
        ("Часовой пояс Timestamp", dtz.key),
        ("Ожидалось точек", len(expected)),
        ("Полных точек", len(complete)),
        ("Пропущенные точки", ", ".join(f"{dt.datetime.fromtimestamp(t, dtz):%H:%M}" for t in missing) or "—"),
        ("Проблемных ячеек", len(issues)),
    ] + [(k, v) for k, v in notes]
    for k, v in rows:
        ch.append([k, v])
        ch.cell(row=ch.max_row, column=1).font = Font(bold=True)
    ch.append([])
    ch.append(["Столбец", "Ожидалось", "Есть значений", "Нет данных", "NaN/Inf", "Вне 0–100 / RAM=0", "Отрицательные", "Сбросы счётчика"])
    hdr = ch.max_row
    per = defaultdict(Counter)
    for _, title, _, prob in issues:
        per[title][prob] += 1
    for c in cols:
        have = sum(1 for ts in expected if ts in c.data)
        p = per[c.title]
        ch.append([c.title, len(expected), have, p["нет данных"], p["NaN/Inf"],
                   p["вне диапазона 0–100"] + p["RAM = 0"], p["отрицательное значение"], p["сброс счётчика"]])
        if sum(p.values()):
            for i in range(1, 9):
                ch.cell(row=ch.max_row, column=i).fill = RED
    ch.append([])
    ch.append(["Timestamp", "Столбец", "Значение", "Проблема"])
    hdr2 = ch.max_row
    for ts, title, v, prob in sorted(issues):
        ch.append([f"{dt.datetime.fromtimestamp(ts, dtz):%Y-%m-%d %H:%M}", title, "" if v is None else str(round(v, 2)), prob])
    for r in (hdr, hdr2):
        for cell in ch[r]:
            cell.font, cell.fill = Font(bold=True), HEAD
    ch.column_dimensions["A"].width = 26
    ch.column_dimensions["B"].width = 60
    for col in "CDEFGH":
        ch.column_dimensions[col].width = 16
    wb.save(path)


def summarize(day, expected, complete, issues, late_note):
    head = f"Отчёт за {day.isoformat()}: {len(complete)}/{len(expected)} точек"
    if not issues and len(complete) == len(expected):
        return head + ", проблем нет"
    cnt = Counter(p for *_, p in issues)
    parts = [head + "; проблемы:"] + [f"• {k}: {v}" for k, v in cnt.most_common()]
    if late_note:
        parts.append(late_note)
    parts.append("Подробности — лист «Проверка».")
    return "\n".join(parts)[:1000]  # лимит подписи Telegram — 1024


def tg_send(cfg, path, caption, attempts=5):
    url = f"https://api.telegram.org/bot{cfg['tg_token']}/sendDocument"
    delay = 10
    for n in range(1, attempts + 1):
        try:
            with open(path, "rb") as f:
                r = requests.post(url, data={"chat_id": cfg["tg_chat"], "caption": caption},
                                  files={"document": (os.path.basename(path), f)}, timeout=60)
            if r.ok and r.json().get("ok"):
                return
            err = f"HTTP {r.status_code}: {r.text[:200]}"
        except requests.RequestException as e:
            err = str(e).replace(cfg["tg_token"], "***")  # токен есть в URL
        log.warning("sendDocument попытка %d/%d: %s", n, attempts, err)
        if n < attempts:
            time.sleep(delay)
            delay = min(delay * 2, 120)
    raise RuntimeError(f"Telegram sendDocument не удался после {attempts} попыток: {err}")


def raise_alert(cfg, text):
    now = dt.datetime.now(dt.timezone.utc)
    alert = [{
        "labels": {"alertname": "DailyReportFailed", "severity": "critical", "host": "rpi1"},
        "annotations": {"summary": "Ошибка ежедневного отчёта", "description": text[:500]},
        "startsAt": now.isoformat(), "endsAt": (now + dt.timedelta(hours=6)).isoformat(),
    }]
    try:
        requests.post(cfg["am_url"].rstrip("/") + "/api/v2/alerts", json=alert, timeout=15).raise_for_status()
    except requests.RequestException as e:
        log.error("не удалось отправить алерт в Alertmanager: %s", e)
    if cfg["hc_url"]:
        try:
            requests.post(cfg["hc_url"].rstrip("/") + "/fail", data=text[:10000].encode(), timeout=15)
        except requests.RequestException:
            pass


def default_day(cfg):
    now = dt.datetime.now(cfg["tz"])
    last_point = dt.datetime.combine(now.date(), dt.time.fromisoformat(cfg["end"]), cfg["tz"]) - dt.timedelta(seconds=STEP)
    return now.date() if now >= last_point else now.date() - dt.timedelta(days=1)


def run(args, cfg):
    day = dt.date.fromisoformat(args.date) if args.date else default_day(cfg)
    tz = cfg["tz"]
    start = dt.datetime.combine(day, dt.time.fromisoformat(cfg["start"]), tz)
    end = dt.datetime.combine(day, dt.time.fromisoformat(cfg["end"]), tz) - dt.timedelta(seconds=STEP)
    expected = list(range(int(start.timestamp()), int(end.timestamp()) + 1, STEP))
    vm = VM(cfg)

    # Ждём, пока vmagent дошлёт буфер (только для «свежего» дня, до wait_max секунд)
    deadline = time.monotonic() + (cfg["wait_max"] if day == dt.datetime.now(tz).date() and not args.no_wait else 0)
    late_note = ""
    while True:
        cols = collect(vm, cfg, int(start.timestamp()), int(end.timestamp()))
        cutoff = int(time.time()) - 90  # задержка поступления данных (latencyOffset VM + интервал сбора)
        issues, cells, complete = validate(cols, expected, cutoff)
        past = [ts for ts in expected if ts <= cutoff]
        lag, pend = buffer_state(vm, cfg)
        log.info("точек %d/%d, проблем %d, lag=%.0fs, буфер=%s", len(complete), len(expected), len(issues), lag, pend)
        buffering = lag > 180 or (not math.isnan(pend) and pend > 1_000_000)
        if len(complete) == len(past) or not buffering:
            break
        if time.monotonic() >= deadline:
            late_note = (f"⚠ Буфер vmagent не опустел за {cfg['wait_max'] // 60} мин "
                         f"(задержка данных {lag:.0f} с) — пропуски могут быть досланы позже, "
                         f"перегенерируйте: report.py --date {day.isoformat()}")
            break
        time.sleep(60)

    notes = [("Сформирован", dt.datetime.now(cfg["dtz"]).strftime("%Y-%m-%d %H:%M:%S ") + cfg["dtz"].key),
             ("Задержка данных сервера, с", round(lag, 0) if not math.isinf(lag) else "нет данных"),
             ("Буфер vmagent (RPI1), байт", "н/д" if math.isnan(pend) else int(pend))]
    if late_note:
        notes.append(("Внимание", late_note))

    os.makedirs(args.out or cfg["out_dir"], exist_ok=True)
    path = os.path.join(args.out or cfg["out_dir"], f"report_{day.isoformat()}.xlsx")
    if len(past) < len(expected):
        notes.append(("Внимание", f"Отчёт сформирован до конца рабочего дня: прошло {len(past)} из {len(expected)} точек"))
    build_xlsx(path, day, cols, expected, cells, issues, complete, notes, tz, cfg["dtz"])
    caption = summarize(day, expected, complete, issues, late_note)
    if len(past) < len(expected):
        caption = caption.replace(f"{len(complete)}/{len(expected)} точек",
                                  f"{len(complete)}/{len(expected)} точек (прошло {len(past)})", 1)
    log.info("файл: %s", path)
    print(caption)

    if args.no_send:
        return
    tg_send(cfg, path, caption)
    log.info("отправлено в Telegram")
    try:
        vm.push_metric(f'srvmon_report_last_success_timestamp_seconds{{job="srvmon-report"}} {int(time.time())}\n')
    except requests.RequestException as e:
        log.warning("не удалось записать метрику успеха: %s", e)
    if cfg["hc_url"]:
        try:
            requests.get(cfg["hc_url"], timeout=15)
        except requests.RequestException as e:
            log.warning("healthchecks ping: %s", e)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--date", help="YYYY-MM-DD (по умолчанию — последний завершившийся рабочий день)")
    ap.add_argument("--no-send", action="store_true", help="только сформировать xlsx")
    ap.add_argument("--no-wait", action="store_true", help="не ждать опустошения буфера vmagent")
    ap.add_argument("--config", default="/etc/srvmon/report.env")
    ap.add_argument("--out", help="каталог для xlsx")
    args = ap.parse_args()
    logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
    cfg = load_config(args.config)
    try:
        run(args, cfg)
    except Exception as e:
        log.exception("отчёт не сформирован")
        if not args.no_send:
            raise_alert(cfg, f"{type(e).__name__}: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
